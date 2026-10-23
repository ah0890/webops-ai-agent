import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.reporting.excel import generate_run_report


@pytest.fixture()
def client(seeded):
    from app.api.main import create_app

    with TestClient(create_app()) as c:
        yield c


def test_overview_reflects_database(client, seeded):
    o = client.get("/api/dashboard/overview").json()
    assert o["websites_monitored"] == 1 and o["products_monitored"] == 2
    assert o["checks_total"] == 4 and o["successful_checks"] == 3 and o["failed_checks"] == 1
    assert o["price_changes_30d"] == 1 and o["price_increases_30d"] == 1
    assert o["last_execution"]["id"] == seeded["r2"].id


def test_products_and_history(client, seeded):
    products = {p["name"]: p for p in client.get("/api/products").json()}
    a = products["Widget A"]["stats"]
    assert a["current_price"] == 110 and a["previous_price"] == 100 and a["percentage_change"] == 10.0
    assert products["Widget B"]["latest_status"] == "FAILED"
    # Failed observations never enter the trusted series.
    b = client.get(f"/api/products/{seeded['p2'].id}/history").json()
    assert len(b["series"]) == 1 and len(b["observations"]) == 2


def test_run_report_json(client, seeded):
    r = client.get(f"/api/runs/{seeded['r2'].id}/report").json()
    s = r["summary"]
    assert (s["products_checked"], s["verified"], s["failed"], s["price_increases"]) == (2, 1, 1, 1)
    assert r["changes"][0]["percentage_change"] == 10.0
    assert r["errors"][0]["status"] == "FAILED"


def test_excel_export_contains_real_rows(client, seeded):
    resp = client.get(f"/api/runs/{seeded['r2'].id}/export")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Executive Summary", "Current Prices", "Price Changes", "Historical Observations",
                             "Availability Changes", "Errors - Failed Checks"]
    changes = list(wb["Price Changes"].iter_rows(min_row=2, values_only=True))
    assert changes[0][1] == "Widget A" and changes[0][7] == 10.0
    assert wb["Historical Observations"].max_row == 5  # header + 4 observations


def test_report_file_is_recorded(seeded, settings):
    path = generate_run_report(seeded["r1"].id, settings, with_ai_summary=False)
    assert path.exists() and path.parent == settings.reports_dir


def test_workflow_create_from_instruction_and_run_conflict(client, seeded):
    r = client.post("/api/workflows", json={"instruction": "Check prices every 6 hours, alert above 7%"})
    assert r.status_code == 201
    wf = r.json()
    assert wf["definition"]["price_change_threshold"] == 7 and wf["schedule_text"] == "Every 6 hours"
    assert wf["source"] == "planner" and wf["next_run_at"]
    assert client.post("/api/workflows", json={}).status_code == 422
    assert client.post("/api/workflows/999/run").status_code == 404


def test_patch_workflow_schedule(client, seeded):
    r = client.patch(f"/api/workflows/{seeded['wf'].id}", json={"schedule": {"kind": "daily", "at": "07:15"}})
    assert r.json()["schedule_text"] == "Daily at 07:15" and r.json()["next_run_at"]


def test_alert_lifecycle_and_errors(client, seeded, session):
    from app.database.models import Alert

    session.add(Alert(run_id=seeded["r2"].id, alert_type="price_increase", severity="medium", reason="test"))
    session.commit()
    aid = client.get("/api/alerts?status=open").json()[0]["id"]
    assert client.patch(f"/api/alerts/{aid}", json={"status": "resolved"}).json()["status"] == "resolved"
    assert client.patch(f"/api/alerts/{aid}", json={"status": "bogus"}).status_code == 422
    assert client.get("/api/errors").json()[0]["extraction_status"] == "FAILED"


def test_import_endpoint(client):
    csv = b"supplier,product_name,product_url\nDelta,Thing,https://delta.test/t\nDelta,,https://delta.test/u\n"
    r = client.post("/api/import/products?dry_run=false", files={"file": ("p.csv", csv, "text/csv")}).json()
    assert (r["imported"], r["invalid"]) == (1, 1)
    bad = client.post("/api/import/products", files={"file": ("p.exe", b"x", "application/octet-stream")})
    assert bad.status_code == 422


def test_api_key_protects_mutations(seeded, monkeypatch):
    from app.api.main import create_app
    from app.config.settings import get_settings

    monkeypatch.setenv("API_KEY", "k-123")
    get_settings.cache_clear()
    with TestClient(create_app()) as c:
        assert c.post("/api/workflows", json={"instruction": "daily"}).status_code == 401
        assert c.post("/api/workflows", json={"instruction": "daily"}, headers={"X-API-Key": "k-123"}).status_code == 201
        assert c.get("/api/workflows").status_code == 200  # reads stay open


def test_screenshot_path_traversal_blocked(client, seeded, session):
    from app.database.models import Screenshot

    s = Screenshot(run_id=seeded["r1"].id, kind="error", path="../../pyproject.toml")
    session.add(s)
    session.commit()
    assert client.get(f"/api/screenshots/{s.id}").status_code == 404


def test_blank_api_key_in_env_means_auth_disabled(monkeypatch):
    from app.config.settings import Settings

    monkeypatch.setenv("API_KEY", "  ")
    monkeypatch.setenv("ALERT_WEBHOOK_URL", "")
    s = Settings(_env_file=None)
    assert s.api_key is None and s.alert_webhook_url is None
