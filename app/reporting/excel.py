"""Excel business report for a run (6 sheets), built only from stored observations."""

from __future__ import annotations

import json
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

from app.agent.llm import LLMUnavailable, get_provider
from app.config.settings import Settings
from app.database.models import Report
from app.database.session import session_scope
from app.services.run_report import build_run_report

log = logging.getLogger(__name__)
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="111827")
UP_FILL = PatternFill("solid", fgColor="FDE2E1")
DOWN_FILL = PatternFill("solid", fgColor="DCFCE7")
STATUS_FILLS = {"VERIFIED": "DCFCE7", "LIKELY_VALID": "E0F2FE", "NEEDS_REVIEW": "FEF3C7", "FAILED": "FDE2E1"}


class _Summary(BaseModel):
    summary: str


def _table(ws: Worksheet, start_row: int, columns: list[tuple[str, str, str | None]], rows: list[dict],
           empty_message: str | None = None) -> int:
    """columns: (header, key, number_format). Returns next free row."""
    for c, (header, _, _) in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=c, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, start=start_row + 1):
        for c, (_, key, fmt) in enumerate(columns, start=1):
            v = row.get(key)
            if isinstance(v, bool):
                v = "yes" if v else "no"
            cell = ws.cell(row=r, column=c, value=v)
            if fmt:
                cell.number_format = fmt
            if key == "status" and v in STATUS_FILLS:
                cell.fill = PatternFill("solid", fgColor=STATUS_FILLS[v])
    end = start_row + len(rows)
    if not rows and empty_message:
        ws.cell(row=start_row + 1, column=1, value=empty_message).font = Font(italic=True, color="6B7280")
        end += 1
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{end}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    for c, (header, key, _) in enumerate(columns, start=1):
        width = max([len(header)] + [len(str(r.get(key) or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(c)].width = min(max(10, width + 2), 60)
    return end + 1


def _pct_rules(ws: Worksheet, col_letter: str, first: int, last: int) -> None:
    if last < first:
        return
    rng = f"{col_letter}{first}:{col_letter}{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=UP_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=DOWN_FILL))


def _ai_summary(data: dict, settings: Settings) -> str | None:
    provider = get_provider(settings)
    if provider is None:
        return None
    facts = {"summary": data["summary"], "top_movers": data["top_movers"][:5], "alerts": data["alerts"][:20]}
    try:
        out = provider.structured(
            "Write a 3-5 sentence executive summary of a price-monitoring run for a procurement manager. Use ONLY "
            "the facts provided; do not invent numbers. Mention data-quality problems if any.",
            json.dumps(facts, default=str), _Summary, max_tokens=1500)
        return out.summary
    except LLMUnavailable as e:
        log.info("AI summary skipped: %s", e)
        return None


def render_workbook(data: dict, ai_summary: str | None = None) -> Workbook:
    wb = Workbook()
    s = data["summary"]

    ws = wb.active
    ws.title = "Executive Summary"
    ws["A1"] = "WebOps AI Agent — Price Monitoring Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "All figures are computed from observations stored for this run. Times are UTC."
    ws["A2"].font = Font(italic=True, color="6B7280")
    kv = [
        ("Workflow", s["workflow"]), ("Run ID", s["run_id"]), ("Status", s["status"]), ("Trigger", s["trigger"]),
        ("Started", s["started_at"]), ("Finished", s["finished_at"]), ("Duration (s)", s["duration_seconds"]),
        ("Websites checked", s["websites_checked"]), ("Products checked", s["products_checked"]),
        ("Verified", s["verified"]), ("Likely valid", s["likely_valid"]), ("Needs review", s["needs_review"]),
        ("Failed", s["failed"]), ("Price increases", s["price_increases"]), ("Price decreases", s["price_decreases"]),
        ("Availability changes", s["availability_changes"]),
        ("Alerts (high / medium / low)", f"{s['alerts_high']} / {s['alerts_medium']} / {s['alerts_low']}"),
    ]
    if s.get("error"):
        kv.append(("Run error", s["error"]))
    for i, (k, v) in enumerate(kv, start=4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    row = 4 + len(kv) + 1
    if ai_summary:
        ws.cell(row=row, column=1, value="AI summary (generated from the figures above)").font = Font(bold=True)
        c = ws.cell(row=row + 1, column=1, value=ai_summary)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=6)
        ws.row_dimensions[row + 1].height = 90
        row += 3
    ws.cell(row=row, column=1, value="Top movers").font = Font(bold=True, size=12)
    if data["top_movers"]:
        end = _table(ws, row + 1, [("Website", "website", None), ("Product", "product", None),
                                   ("Previous", "previous_price", "#,##0.00"), ("Current", "current_price", "#,##0.00"),
                                   ("Change %", "percentage_change", "0.00")], data["top_movers"])
        _pct_rules(ws, "E", row + 2, end - 1)
    else:
        ws.cell(row=row + 1, column=1, value="No price changes versus previous trusted observations in this run.")
    ws.freeze_panes = None
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 40

    ws = wb.create_sheet("Current Prices")
    end = _table(ws, 1, [
        ("Website", "website", None), ("Product", "product", None), ("SKU", "sku", None),
        ("Price", "price", "#,##0.00"), ("Currency", "currency", None), ("Availability", "availability", None),
        ("Data quality", "status", None), ("Method", "method", None), ("Observed (UTC)", "observed_at", None),
        ("Previous", "previous_price", "#,##0.00"), ("Abs change", "absolute_change", "#,##0.00"),
        ("Change %", "percentage_change", "0.00"), ("7d %", "change_7d_pct", "0.00"), ("30d %", "change_30d_pct", "0.00"),
        ("Lowest", "lowest_price", "#,##0.00"), ("Highest", "highest_price", "#,##0.00"),
        ("Average", "average_price", "#,##0.00"), ("Trend", "trend", None), ("# obs", "observations", None),
        ("URL", "url", None),
    ], data["current"])
    _pct_rules(ws, "L", 2, end - 1)

    ws = wb.create_sheet("Price Changes")
    end = _table(ws, 1, [
        ("Website", "website", None), ("Product", "product", None), ("SKU", "sku", None),
        ("Previous", "previous_price", "#,##0.00"), ("Current", "current_price", "#,##0.00"), ("Currency", "currency", None),
        ("Abs change", "absolute_change", "#,##0.00"), ("Change %", "percentage_change", "0.00"),
        ("Direction", "direction", None), ("Threshold %", "threshold", None), ("Alerted", "alerted", None),
        ("Previous observed", "previous_observed_at", None), ("Observed", "observed_at", None),
    ], data["changes"], empty_message="No price changes versus the previous trusted observation were detected in this run.")
    _pct_rules(ws, "H", 2, end - 1)

    ws = wb.create_sheet("Historical Observations")
    _table(ws, 1, [
        ("Website", "website", None), ("Product", "product", None), ("Observed (UTC)", "observed_at", None),
        ("Price", "price", "#,##0.00"), ("Currency", "currency", None), ("Availability", "availability", None),
        ("Data quality", "status", None), ("Method", "method", None), ("Run", "run_id", None), ("URL", "url", None),
    ], data["history"])

    ws = wb.create_sheet("Availability Changes")
    _table(ws, 1, [
        ("Website", "website", None), ("Product", "product", None), ("Previous", "previous", None),
        ("Current", "current", None), ("Previous observed", "previous_observed_at", None),
        ("Observed", "observed_at", None), ("URL", "url", None),
    ], data["availability"], empty_message="No availability changes were detected in this run.")

    ws = wb.create_sheet("Errors - Failed Checks")
    end = _table(ws, 1, [
        ("Website", "website", None), ("Product", "product", None), ("Data quality", "status", None),
        ("Issues", "issues", None), ("Observed (UTC)", "observed_at", None), ("Screenshot", "screenshot", None),
        ("URL", "url", None),
    ], data["errors"], empty_message="No failed or review-flagged checks in this run.")
    if data["failed_steps"]:
        ws.cell(row=end + 1, column=1, value="Failed execution steps").font = Font(bold=True)
        _table(ws, end + 2, [("Stage", "stage", None), ("Product ID", "product_id", None),
                             ("Attempt", "attempt", None), ("Message", "message", None), ("At", "at", None)],
               data["failed_steps"])
        ws.freeze_panes = "A2"
    return wb


def generate_run_report(run_id: int, settings: Settings, with_ai_summary: bool = True) -> Path:
    with session_scope() as s:
        data = build_run_report(s, run_id)
    ai = _ai_summary(data, settings) if with_ai_summary else None
    wb = render_workbook(data, ai)
    settings.reports_dir.mkdir(parents=True, exist_ok=True)
    started = (data["summary"]["started_at"] or "").replace(":", "").replace("-", "").replace(" ", "_")
    path = settings.reports_dir / f"price_report_run{run_id}_{started}Z.xlsx"
    wb.save(path)
    with session_scope() as s:
        s.add(Report(run_id=run_id, path=path.name, summary={**data["summary"], "ai_summary": ai}))
        s.commit()
    return path
